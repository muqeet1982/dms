"""
Product / price-list importer.

Understands two layouts:

1. The flat CSV/XLSX template (one header row, one product per row).
2. A supplier "Packing & Price List" workbook as distributors receive it:

        Keune Pakistan
        Price List Jul-2026
        2026-09-15
        TINTA COLOR                                  <- section row = category
        S.No | Code | Product Name | UNIT | Packing | D.P | T.P | R.P
        1    | 62001| NO. 1 BLACK - 60 ML | 60ml | 24 | 1130 | 1255 | 1495
        ...
        COLOR CHAMELEON                              <- next category
        ...

Column headers are matched by alias (Code/SKU, Product Name/Name, UNIT/Size, Packing/Carton,
D.P/Cost, T.P/Sale, R.P/MRP …) so the same workbook can be re-uploaded whenever prices change:
products are matched by code and updated in place, new codes are created.
"""

import csv
import io
import re
from datetime import date, datetime
from decimal import Decimal, InvalidOperation

ALIASES = {
    "sku": ["sku", "code", "item", "item code", "item no", "item#", "product code", "itemcode", "article", "article no", "part no"],
    "name": ["name", "product name", "product", "item name", "item description", "description of goods", "particulars"],
    "unit_size": ["unit", "size", "unit size", "pack size", "packing size", "unit_size", "volume", "weight"],
    "packing": ["packing", "carton", "ctn", "carton qty", "carton_qty", "units per carton", "pcs per carton", "pcs/ctn", "pack of", "carton size", "packing qty"],
    "cost_price": ["dp", "d p", "distributor price", "dist price", "cost", "cost price", "cost_price", "net price"],
    "sale_price": ["tp", "t p", "trade price", "sale price", "sale_price", "selling price", "wholesale price", "price"],
    "retail_price": ["rp", "r p", "retail price", "retail_price", "mrp", "retail", "consumer price"],
    "category": ["category", "group", "section", "product group", "range", "line"],
    "brand": ["brand", "make"],
    "supplier": ["supplier", "vendor", "manufacturer", "company"],
    "barcode": ["barcode", "ean", "ean13", "gtin"],
    "upc": ["upc", "unit upc", "unit barcode", "upc code", "piece barcode"],
    "box_upc": ["box upc", "box barcode", "carton upc", "carton barcode", "outer barcode", "case upc"],
    "item_code": ["item code", "itemcode", "item no", "item#", "product code", "article", "article no", "part no", "supplier code"],
    "uom": ["uom", "base uom", "base unit"],
    "tax_class": ["tax class", "tax_class"],
    "min_stock": ["min stock", "min_stock", "minimum stock"],
    "max_stock": ["max stock", "max_stock", "maximum stock"],
    "reorder_qty": ["reorder qty", "reorder_qty", "reorder"],
    "batch_tracked": ["batch tracked", "batch_tracked", "batch"],
    "expiry_tracked": ["expiry tracked", "expiry_tracked", "expiry"],
    "box_qty": ["box qty", "box_qty", "box", "inner", "inner pack"],
    "description": ["description", "notes", "remarks"],
    "status": ["status"],
    "location": ["location", "bin", "bin code", "rack", "shelf", "storage", "location code", "store location", "place"],
    "opening_qty": ["opening qty", "opening stock", "opening quantity", "opening", "opening balance qty", "stock qty", "qty on hand", "on hand", "stock"],
    "opening_unit_cost": ["opening cost", "opening unit cost", "opening rate"],
    # order sheets (purchase orders)
    "qty": ["qty", "quantity", "order qty", "order quantity", "order", "demand", "required", "req qty", "pcs", "pieces", "units", "qty pcs", "order pcs"],
    "cartons": ["cartons", "ctns", "no of cartons", "order ctn", "order cartons", "qty ctn"],
    "unit_price": ["rate", "unit price", "unit_price", "purchase rate", "po rate", "cost rate", "purchase price"],
}
_ALIAS_INDEX = {alias: key for key, aliases in ALIASES.items() for alias in aliases}
SKIP_HEADERS = {"sno", "s no", "sr", "sr no", "serial", "#", "no"}


def norm_header(v) -> str:
    """'D.P' → 'dp', ' Product  Name ' → 'product name', 'Carton_Qty' → 'carton_qty'"""
    s = re.sub("[.\u00a0]", "", str(v or "")).strip().lower()  # drop dots and non-breaking spaces
    return re.sub(r"\s+", " ", s)


def canonical(header) -> str | None:
    h = norm_header(header)
    if not h or h in SKIP_HEADERS:
        return None
    return _ALIAS_INDEX.get(h) or _ALIAS_INDEX.get(h.replace("_", " ")) or _ALIAS_INDEX.get(h.replace(" ", "_"))


def cell_text(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    if isinstance(v, (datetime, date)):
        return v.isoformat()
    return str(v).strip()


def to_decimal(v, default=None):
    t = cell_text(v).replace(",", "").replace("Rs", "").replace("PKR", "").strip()
    if t in ("", "-", "—"):
        return default
    try:
        return Decimal(t)
    except InvalidOperation as e:
        raise ValueError(f"'{v}' is not a number") from e


def to_int(v, default=None):
    d = to_decimal(v, None)
    return default if d is None else int(d)


def to_bool(v) -> bool:
    return cell_text(v).lower() in ("1", "true", "yes", "y", "on")


def titled(s: str) -> str:
    """Section headings arrive in CAPS; store them as 'Tinta Color', 'Fix & Finish', 'So Pure Treatments'."""
    s = re.sub(r"\s+", " ", (s or "").strip())
    return s.title() if s.isupper() else s


class SheetReader:
    """Turns a CSV or XLSX upload into (rows, meta). Each row is {canonical_key: raw_value}; meta carries title rows and a date."""

    @staticmethod
    def read(file, filename: str):
        name = (filename or "").lower()
        if name.endswith((".xlsx", ".xlsm", ".xltx")):
            return SheetReader._read_xlsx(file)
        if name.endswith(".xls"):
            raise ValueError("Legacy .xls workbooks are not supported — save the file as .xlsx (Excel 2007+) and upload again.")
        return SheetReader._read_csv(file)

    @staticmethod
    def _read_csv(file):
        text = file.read()
        if isinstance(text, bytes):
            text = text.decode("utf-8-sig", errors="replace")
        grid = [list(r) for r in csv.reader(io.StringIO(text))]
        return SheetReader._parse_grid(grid)

    @staticmethod
    def _read_xlsx(file):
        import openpyxl

        wb = openpyxl.load_workbook(file, data_only=True, read_only=True)
        best, best_meta = [], {}
        for ws in wb.worksheets:
            grid = [list(r) for r in ws.iter_rows(values_only=True)]
            rows, meta = SheetReader._parse_grid(grid)
            if len(rows) > len(best):
                best, best_meta = rows, {**meta, "sheet": ws.title}
        return best, best_meta

    MONTHS = {m: i for i, m in enumerate(["jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep", "oct", "nov", "dec"], start=1)}

    @staticmethod
    def _parse_date(text):
        """'21-Sep-2026', '21/09/2026', '2026-09-21' → date; anything else → None."""
        t = (text or "").strip()
        m = re.fullmatch(r"(\d{1,2})[-/ ]([A-Za-z]{3,9})[-/ ](\d{2,4})", t)
        if m and m.group(2)[:3].lower() in SheetReader.MONTHS:
            d, mo, y = int(m.group(1)), SheetReader.MONTHS[m.group(2)[:3].lower()], int(m.group(3))
            return date(y + 2000 if y < 100 else y, mo, d)
        m = re.fullmatch(r"(\d{1,2})[-/.](\d{1,2})[-/.](\d{2,4})", t)
        if m:
            d, mo, y = (int(x) for x in m.groups())
            return date(y + 2000 if y < 100 else y, mo, d)
        m = re.fullmatch(r"(\d{4})-(\d{1,2})-(\d{1,2})", t)
        if m:
            y, mo, d = (int(x) for x in m.groups())
            return date(y, mo, d)
        return None

    @staticmethod
    def _parse_grid(grid):
        meta = {"title_rows": [], "date": None, "sections": [], "labels": {}}
        header_idx, columns = None, {}
        for i, raw in enumerate(grid):
            found = {}
            for ci, cell in enumerate(raw):
                key = canonical(cell)
                if key and key not in found:
                    found[key] = ci
            if "sku" in found and "name" in found:
                header_idx, columns = i, found
                break
            # preamble: remember free text (brand / title), any date, and "LABEL: value" pairs (S.O. NUMBER: 51145)
            texts = [(ci, cell) for ci, cell in enumerate(raw) if cell is not None and cell_text(cell) != ""]
            for pos, (ci, cell) in enumerate(texts):
                if isinstance(cell, (datetime, date)):
                    meta["date"] = meta["date"] or (cell.date() if isinstance(cell, datetime) else cell)
                    continue
                if not isinstance(cell, str):
                    continue
                text = cell.strip()
                parsed = SheetReader._parse_date(text)
                if parsed and not meta["date"]:
                    meta["date"] = parsed
                    continue
                m = re.match(r"^(.{2,40}?)\s*[:：]\s*(.*)$", text)  # "S.O. NUMBER: 51145" or "S.O. DATE:" + next cell
                if m:
                    label = norm_header(m.group(1))
                    value = m.group(2).strip() or (cell_text(texts[pos + 1][1]) if pos + 1 < len(texts) else "")
                    if label and value:
                        meta["labels"][label] = value
                        continue
                meta["title_rows"].append(text)
        if header_idx is None:
            raise ValueError("Could not find the header row — the sheet needs a 'Code' (or SKU) column and a 'Product Name' column.")
        name_col, sku_col = columns["name"], columns["sku"]
        price_cols = [columns[k] for k in ("cost_price", "sale_price", "retail_price") if k in columns]
        if "qty" not in columns and "cartons" not in columns:
            # Order sheets often carry the quantity in a column with no heading: take the right-most header-less
            # column that holds numbers on product rows (ignoring the S.No column and everything already mapped).
            header = grid[header_idx]
            known = set(columns.values()) | {ci for ci, c in enumerate(header) if norm_header(c) in SKIP_HEADERS}
            counts = {}
            for raw in grid[header_idx + 1 :]:
                if sku_col < len(raw) and cell_text(raw[sku_col]) and name_col < len(raw) and cell_text(raw[name_col]):
                    for ci, cell in enumerate(raw):
                        if ci not in known and (ci >= len(header) or cell_text(header[ci]) == "") and isinstance(cell, (int, float)) and not isinstance(cell, bool):
                            counts[ci] = counts.get(ci, 0) + 1
            if counts:
                columns["qty"] = max(counts)  # right-most
                meta["qty_column_guessed"] = True
        rows, section = [], None
        # A lone text cell on the row right above the header is the first section (e.g. "TINTA COLOR"), not a title.
        if header_idx > 0:
            above = [cell_text(c) for c in grid[header_idx - 1] if isinstance(c, str) and cell_text(c)]
            # section headings are printed in CAPS ("TINTA COLOR"); a mixed-case line ("Lahore order") is a title
            if len(above) == 1 and above[0].isupper() and meta["title_rows"] and meta["title_rows"][-1] == above[0]:
                section = titled(meta["title_rows"].pop())
                meta["sections"].append(section)
        for i, raw in enumerate(grid[header_idx + 1 :], start=header_idx + 2):
            if all(cell_text(c) == "" for c in raw):
                continue
            sku, name = cell_text(raw[sku_col]) if sku_col < len(raw) else "", cell_text(raw[name_col]) if name_col < len(raw) else ""
            has_price = any(ci < len(raw) and cell_text(raw[ci]) not in ("", "0") for ci in price_cols)
            if (not sku or not name) and not has_price:
                # a section heading (category): a lone text cell (in the code column, the name column or elsewhere), no prices
                texts = [cell_text(c) for c in raw if isinstance(c, str) and cell_text(c)]
                if texts:
                    section = titled(texts[0])
                    if section not in meta["sections"]:
                        meta["sections"].append(section)
                continue
            row = {k: raw[ci] for k, ci in columns.items() if ci < len(raw) and cell_text(raw[ci]) != ""}
            if section and cell_text(row.get("category")) == "":
                row["category"] = section
            row["_row"] = i
            rows.append(row)
        return rows, meta


def read_flat(file, filename: str):
    """Generic reader for the other import types: first non-empty row is the header; keys are lower-cased headers."""
    name = (filename or "").lower()
    if name.endswith((".xlsx", ".xlsm", ".xltx")):
        import openpyxl

        ws = openpyxl.load_workbook(file, data_only=True, read_only=True).worksheets[0]
        grid = [list(r) for r in ws.iter_rows(values_only=True)]
    else:
        text = file.read()
        if isinstance(text, bytes):
            text = text.decode("utf-8-sig", errors="replace")
        grid = [list(r) for r in csv.reader(io.StringIO(text))]
    rows, headers = [], None
    for i, raw in enumerate(grid, start=1):
        if all(cell_text(c) == "" for c in raw):
            continue
        if headers is None:
            headers = [norm_header(c).replace(" ", "_") for c in raw]
            continue
        row = {h: cell_text(raw[ci]) for ci, h in enumerate(headers) if h and ci < len(raw)}
        row["_row"] = i
        rows.append(row)
    return rows


class ProductImporter:
    """Upserts one product per row. Matching key: SKU/code (case-insensitive)."""

    def __init__(self, options: dict, meta: dict):
        from apps.catalog.models import Brand, UnitOfMeasure
        from apps.inventory.models import Warehouse
        from apps.partners.models import Supplier

        self.options = options or {}
        self.meta = meta or {}
        self.summary = {"created": 0, "updated": 0, "categories_created": 0, "locations_created": 0, "opening_stock_posted": 0}
        self.update_existing = str(self.options.get("update_existing", "true")).lower() not in ("false", "0", "no")
        self.brand = self._resolve(Brand, self.options.get("brand"), create=True)
        self.supplier = self._resolve(Supplier, self.options.get("supplier"), create=False)
        eff = self.options.get("price_effective_date") or self.meta.get("date")
        self.price_date = date.fromisoformat(str(eff)[:10]) if eff else None
        self.default_uom, _ = UnitOfMeasure.objects.get_or_create(code="PC", defaults={"name": "Piece"})
        wh = self.options.get("warehouse")
        self.warehouse = (Warehouse.objects.filter(pk=wh).first() if wh and re.fullmatch(r"[0-9a-f-]{36}", str(wh)) else Warehouse.objects.filter(code__iexact=str(wh or "")).first()) if wh else None
        if self.warehouse is None:  # locations and opening stock land in the main warehouse unless one is chosen
            self.warehouse = Warehouse.objects.filter(is_active=True).order_by("-is_default", "name").first()
        self._cat_cache, self._bin_cache = {}, {}

    @staticmethod
    def _resolve(model, value, create):
        """Option values may be a UUID (from the UI selects) or a name (from a CSV/API call)."""
        if value in (None, "", "null"):
            return None
        value = str(value).strip()
        obj = model.objects.filter(pk=value).first() if re.fullmatch(r"[0-9a-f-]{36}", value) else None
        if obj is None:
            obj = model.objects.filter(name__iexact=value).first()
        if obj is None and create:
            obj = model.objects.create(name=value)
        if obj is None:
            raise ValueError(f"{model.__name__} '{value}' not found")
        return obj

    def category_for(self, name):
        from apps.catalog.models import Category

        name = titled(cell_text(name))[:100]
        if not name:
            return None
        if name not in self._cat_cache:
            cat = Category.objects.filter(name__iexact=name, parent=None).first()
            if cat is None:
                cat = Category.objects.create(name=name, parent=None)
                self.summary["categories_created"] += 1
            self._cat_cache[name] = cat
        return self._cat_cache[name]

    def location_for(self, code):
        """A location (warehouse bin) per code in the sheet — created on first use so sheets can carry rack codes."""
        from apps.inventory.models import Bin

        code = cell_text(code)[:30]
        if not code or self.warehouse is None:
            return None
        key = code.lower()
        if key not in self._bin_cache:
            b = Bin.objects.filter(warehouse=self.warehouse, code__iexact=code).first()
            if b is None:
                b = Bin.objects.create(warehouse=self.warehouse, code=code)
                self.summary["locations_created"] += 1
            self._bin_cache[key] = b
        return self._bin_cache[key]

    def import_row(self, r: dict):
        from apps.catalog.models import Barcode, Brand, Packaging, Product, TaxClass, UnitOfMeasure
        from apps.catalog.services import ProductService
        from apps.core.tenancy import get_current_tenant
        from apps.partners.models import Supplier

        sku = cell_text(r.get("sku"))[:60]
        name = cell_text(r.get("name"))[:200]
        if not sku:
            raise ValueError("Code/SKU is empty")
        if not name:
            raise ValueError("Product name is empty")
        brand = self.brand or (Brand.objects.get_or_create(name=cell_text(r["brand"])[:100])[0] if cell_text(r.get("brand")) else None)
        supplier = self.supplier or (Supplier.objects.filter(name__iexact=cell_text(r["supplier"])).first() if cell_text(r.get("supplier")) else None)
        category = self.category_for(r["category"]) if cell_text(r.get("category")) else None
        tax_class = TaxClass.objects.filter(name__iexact=cell_text(r["tax_class"])).first() if cell_text(r.get("tax_class")) else None
        uom = UnitOfMeasure.objects.get_or_create(code=cell_text(r["uom"]).upper()[:10], defaults={"name": cell_text(r["uom"])})[0] if cell_text(r.get("uom")) else self.default_uom
        location = self.location_for(r.get("location")) if cell_text(r.get("location")) else None

        fields = {}
        if location is not None:
            fields["location"] = location
        if "unit_size" in r:
            fields["unit_size"] = cell_text(r["unit_size"])[:30]
        if "packing" in r:
            fields["packing"] = max(1, to_int(r["packing"], 1))
        for f in ("cost_price", "sale_price", "retail_price"):
            if f in r:
                fields[f] = to_decimal(r[f], Decimal("0"))
        for f in ("min_stock", "max_stock", "reorder_qty"):
            if f in r:
                fields[f] = to_decimal(r[f], Decimal("0"))
        if "batch_tracked" in r:
            fields["is_batch_tracked"] = to_bool(r["batch_tracked"])
        if "expiry_tracked" in r:
            fields["is_expiry_tracked"] = to_bool(r["expiry_tracked"])
        if "description" in r:
            fields["description"] = cell_text(r["description"])
        if cell_text(r.get("status")).lower() in ("active", "inactive", "archived"):
            fields["status"] = cell_text(r["status"]).lower()
        if self.price_date and any(k in fields for k in ProductService.PRICE_FIELDS):
            fields["price_effective_date"] = self.price_date
        # the code column on a price list is the supplier's item code; keep it on the product as well as the key
        for code_field in ("item_code", "upc", "box_upc"):
            value = cell_text(r.get(code_field))[:64]
            if value:
                fields[code_field] = value
        fields.setdefault("item_code", sku[:60])

        product = Product.objects.filter(sku__iexact=sku).select_related("base_uom").first()
        if product is not None:
            if not self.update_existing:
                raise ValueError(f"Product {sku} already exists")
            before = {f: getattr(product, f) for f in (*ProductService.PRICE_FIELDS, "packing", "upc", "box_upc")}
            product.name = name
            for k, v in fields.items():
                setattr(product, k, v)
            if brand is not None:
                product.brand = brand
            if supplier is not None:
                product.supplier = supplier
            if category is not None:
                product.category = category
            if tax_class is not None:
                product.tax_class = tax_class
            if product.status == "archived":
                product.status = "active"
            product.save()
            ProductService.after_update(product, before)
            self.summary["updated"] += 1
            variant = product.variants.filter(is_default=True).first() or product.variants.first()
            if variant is not None and variant.name != name and variant.sku.lower() == sku.lower():
                variant.name = name
                variant.save(update_fields=["name"])
        else:
            data = {"sku": sku, "name": name, "brand": brand, "category": category, "supplier": supplier, "base_uom": uom, "tax_class": tax_class, **fields}
            product = ProductService.create_product(get_current_tenant(), data)
            self.summary["created"] += 1
            variant = product.variants.first()

        opening = to_decimal(r.get("opening_qty"), None)
        if opening is None and "qty" in r:  # a plain "Qty"/"Stock" column on a product sheet is the opening quantity
            opening = to_decimal(r.get("qty"), None)
        if opening and opening > 0:
            from apps.catalog.services import ProductService as _PS
            from apps.inventory.models import StockMovement

            already = StockMovement.objects.filter(variant__product=product, movement_type="opening").exists()
            if not already:
                _PS.post_opening_stock(product, opening, warehouse=self.warehouse, unit_cost=to_decimal(r.get("opening_unit_cost"), None), note="Opening stock (import)")
                self.summary["opening_stock_posted"] += 1

        if variant is not None and cell_text(r.get("barcode")):
            code = cell_text(r["barcode"])[:64]
            Barcode.objects.get_or_create(code=code, defaults={"variant": variant, "type": "EAN13" if len(code) == 13 and code.isdigit() else "CODE128", "is_primary": not variant.barcodes.exists()})
        if variant is not None and to_int(r.get("box_qty"), 0) > 1:
            box, _ = UnitOfMeasure.objects.get_or_create(code="BOX", defaults={"name": "Box"})
            Packaging.objects.update_or_create(variant=variant, name="box", defaults={"uom": box, "qty_in_base": Decimal(to_int(r["box_qty"]))})
        return product


class PurchaseOrderImporter:
    """
    One supplier order sheet → one draft purchase order. Same layout as the price list (Code / Product Name / UNIT …)
    plus a quantity column (headed "Qty"/"Order" … or the unlabeled numeric column at the right). Rows without a
    quantity are skipped; unknown codes are reported per row and the PO is still created from the rest.
    """

    def __init__(self, options: dict, meta: dict, filename: str = ""):
        from apps.inventory.models import Warehouse
        from apps.partners.models import Supplier

        self.options, self.meta, self.filename = options or {}, meta or {}, filename or ""
        self.summary = {"lines": 0, "skipped_no_qty": 0, "unknown_codes": 0}
        self.lines, self.po, self.skipped_rows = [], None, []
        titles = self.meta.get("title_rows") or []
        self.supplier = ProductImporter._resolve(Supplier, self.options.get("supplier"), create=False) if self.options.get("supplier") else None
        if self.supplier is None:
            for t in titles:  # "Keune Pakistan" on the first row of the sheet
                self.supplier = Supplier.objects.filter(name__iexact=t.strip()).first()
                if self.supplier:
                    break
        if self.supplier is None:
            raise ValueError("Supplier not found — choose the supplier for this order (the sheet title did not match a supplier name).")
        wh_opt = self.options.get("warehouse")
        self.warehouse = (Warehouse.objects.filter(pk=wh_opt).first() if wh_opt and re.fullmatch(r"[0-9a-f-]{36}", str(wh_opt)) else Warehouse.objects.filter(code__iexact=str(wh_opt or "")).first()) if wh_opt else None
        if self.warehouse is None:
            self.warehouse = Warehouse.objects.filter(is_active=True).order_by("-is_default", "name").first()
        if self.warehouse is None:
            raise ValueError("No warehouse — create a warehouse under Settings → Warehouses first.")
        labels = self.meta.get("labels", {})
        # A supplier's own sales-order print carries its document number and date ("S.O. NUMBER: 51145")
        self.supplier_ref = (self.options.get("supplier_ref") or labels.get("so number") or labels.get("sales order number") or labels.get("order number") or labels.get("invoice number") or "").strip()[:60]
        self.customer_label = (labels.get("customer name") or "").strip()
        d = self.options.get("date") or self.meta.get("date")
        if not d:
            m = re.search(r"(\d{1,2})-(\d{1,2})-(\d{2,4})", self.filename)  # LHR-Order-20-9-26.xlsx
            if m:
                dd, mm, yy = (int(x) for x in m.groups())
                d = date(yy + 2000 if yy < 100 else yy, mm, dd)
        self.date = date.fromisoformat(str(d)[:10]) if d else date.today()
        ref_from_labels = (labels.get("po number") or "").strip() or (f"S.O. {self.supplier_ref}" + (f" — {self.customer_label}" if self.customer_label else "") if self.supplier_ref else "")
        title_ref = next((t for t in titles[1:] if t.lower() not in ("sales order", "purchase order")), titles[0] if titles else "")
        self.reference = (self.options.get("reference") or "").strip() or ref_from_labels or title_ref or re.sub(r"\.[A-Za-z0-9]+$", "", self.filename.rsplit("/", 1)[-1])
        self.qty_unit = "cartons" if str(self.options.get("qty_unit", "pieces")).lower().startswith("carton") else "pieces"
        self.expected_date = date.fromisoformat(str(self.options["expected_date"])[:10]) if self.options.get("expected_date") else None
        self.notes = (self.options.get("notes") or "").strip()

    def location_for(self, code):
        """A location (warehouse bin) per code in the sheet — created on first use so sheets can carry rack codes."""
        from apps.inventory.models import Bin

        code = cell_text(code)[:30]
        if not code or self.warehouse is None:
            return None
        key = code.lower()
        if key not in self._bin_cache:
            b = Bin.objects.filter(warehouse=self.warehouse, code__iexact=code).first()
            if b is None:
                b = Bin.objects.create(warehouse=self.warehouse, code=code)
                self.summary["locations_created"] += 1
            self._bin_cache[key] = b
        return self._bin_cache[key]

    def import_row(self, r: dict):
        from apps.catalog.models import ProductVariant

        qty = to_decimal(r.get("qty"), None)
        cartons = to_decimal(r.get("cartons"), None)
        if (qty is None or qty <= 0) and (cartons is None or cartons <= 0):
            self.summary["skipped_no_qty"] += 1
            # the upload reports these back on screen, so which lines were left off no longer has to be written
            # into the order's notes by hand
            if cell_text(r.get("sku")) and len(self.skipped_rows) < 50:
                self.skipped_rows.append({"sku": cell_text(r.get("sku"))[:40], "name": cell_text(r.get("name"))[:60]})
            return None
        sku = cell_text(r.get("sku"))
        variant = ProductVariant.objects.filter(sku__iexact=sku, is_active=True).select_related("product").first()
        if variant is None:
            self.summary["unknown_codes"] += 1
            raise ValueError(f"Code {sku} ({cell_text(r.get('name'))[:40]}) is not in the product list — import the price list first")
        packing = Decimal(variant.product.packing or 1)
        packaging = None
        if cartons is not None and cartons > 0:
            packaging = variant.packagings.filter(name__iexact="carton").first()
            line_qty = cartons if packaging else cartons * packing
        elif self.qty_unit == "cartons":
            packaging = variant.packagings.filter(name__iexact="carton").first()
            line_qty = qty if packaging else qty * packing
        else:
            line_qty = qty
        price = to_decimal(r.get("unit_price"), None)
        if price is None:
            price = to_decimal(r.get("cost_price"), None)
        if price is None:
            price = variant.cost_price
        self.lines.append({"variant": variant, "packaging": packaging, "qty": line_qty, "unit_price": price, "discount_pct": to_decimal(r.get("discount_pct"), Decimal("0")), "price_per_packaging": False})
        return variant

    def finish(self, job=None):
        from django.core.files.base import ContentFile
        from django.db import transaction
        from django.db.models import Sum

        from apps.core.models import Attachment
        from apps.purchasing.services import PurchaseOrderService

        if not self.lines:
            raise ValueError("No order lines — the sheet has no quantities (or no product codes matched).")
        with transaction.atomic():
            # Notes stay the buyer's own: what the upload did, and the rows it could not place, are reported from
            # the import job on the screen that uploaded the sheet — and the sheet itself is attached below.
            notes = self.notes
            if self.supplier_ref and not str(self.options.get("allow_duplicate", "")).lower() in ("1", "true", "yes"):
                from apps.purchasing.models import PurchaseOrder

                existing = PurchaseOrder.objects.filter(supplier=self.supplier, supplier_ref=self.supplier_ref).exclude(status="cancelled").first()
                if existing is not None:
                    raise ValueError(f"{self.supplier.name} document {self.supplier_ref} was already imported as {existing.number} ({existing.status}). Tick 'Import anyway' to create a second purchase order.")
            self.po = PurchaseOrderService.create(supplier=self.supplier, warehouse=self.warehouse, lines=self.lines, date=self.date, expected_date=self.expected_date, notes=notes, reference=self.reference[:120], supplier_ref=self.supplier_ref)
            if job is not None and job.file:
                job.file.open("rb")
                try:
                    data = job.file.read()
                finally:
                    job.file.close()
                name = self.filename.rsplit("/", 1)[-1] or "order.xlsx"
                att = Attachment(model="PurchaseOrder", object_id=str(self.po.id), name=name, size=len(data), kind="document", uploaded_by=job.started_by, content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet" if name.lower().endswith(".xlsx") else "text/csv")
                att.file.save(name, ContentFile(data), save=True)
        self.summary.update({"units": str(self.po.lines.aggregate(n=Sum("qty_ordered"))["n"] or 0), "skipped_rows": self.skipped_rows, "lines": len(self.lines), "po_id": str(self.po.id), "po_number": self.po.number, "supplier_ref": self.supplier_ref, "customer_label": self.customer_label, "total": str(self.po.total), "supplier": self.supplier.name, "warehouse": self.warehouse.name, "date": self.date.isoformat(), "reference": self.reference, "qty_unit": self.qty_unit})
        return self.po
