"""Importing a supplier 'Packing & Price List' workbook (Code / Product Name / UNIT / Packing / D.P / T.P / R.P with section rows)."""
import io
from datetime import date, datetime
from decimal import Decimal

import openpyxl
import pytest
from django.core.files.uploadedfile import SimpleUploadedFile

from apps.core.tenancy import tenant_context

pytestmark = pytest.mark.django_db


def price_list_xlsx(prices=("1130", "1255", "1495")):
    """Mimics the real file: title rows, a date, section headings in the Code column, header row, blank spacer column."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Stock_Report"
    dp, tp, rp = prices
    rows = [
        [None, "Keune Pakistan"],
        [None, "Price List Jul-2026"],
        [None, datetime(2026, 9, 15)],
        [None, "TINTA COLOR"],
        [None, "S.No", "Code", "Product Name", None, None, "UNIT", None, "Packing", "D.P", "T.P", "R.P"],
        [None, 1, 62001, "NO. 1 BLACK - 60 ML", None, None, "60ml", None, 24, int(dp), int(tp), int(rp)],
        [None, 2, "62300UC", "NO. 3.00 DARK BROWN (UC) - 60 ML", None, None, "60ml", None, 24, 1166, 1295, 1540],
        [None, None, "COLOR CHAMELEON"],
        [None, 3, 12110, "COLOR CHAMELEON YELLOW", None, None, "60ml", None, 24, 1431, 1590, 1890],
        [None, None, "DANDRUFF DETOX", None, None, None, None, None, None, 0, 0, None],
        [None, 4, 21500, "DANDRUFF DETOX PURIFYING SHAMPOO - 300 ML", None, None, "300ml", None, 12, 3001.5, 3335, 4170],
        [None, None, "ACCESSORIES"],
        [None, 5, 53560, "KEUNE TIMER", None, None, "pcs", None, 1, 405, 450, 540],
    ]
    for r in rows:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def run_import(api, content, options=None, name="Packing & Price List 15-09-26.xlsx"):
    data = {"type": "products", "file": SimpleUploadedFile(name, content, content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    if options:
        import json

        data["options"] = json.dumps(options)
    r = api.post("/api/v1/imports", data, format="multipart")
    assert r.status_code == 201, r.content
    return r.data


def test_price_list_workbook_creates_products_with_categories_and_prices(api, tenant):
    from apps.catalog.models import Category, Product

    job = run_import(api, price_list_xlsx(), {"brand": "Keune"})
    assert job["status"] == "done" and job["rows_failed"] == 0, job["errors"]
    assert job["rows_total"] == 5 and job["summary"]["created"] == 5 and job["summary"]["categories_created"] == 4
    assert job["summary"]["price_list_date"] == "2026-09-15"
    with tenant_context(tenant, tenant.owner):
        p = Product.objects.get(sku="62001")
        assert (p.name, p.unit_size, p.packing) == ("NO. 1 BLACK - 60 ML", "60ml", 24)
        assert (p.cost_price, p.sale_price, p.retail_price) == (Decimal("1130"), Decimal("1255"), Decimal("1495"))
        assert p.brand.name == "Keune" and p.category.name == "Tinta Color" and p.price_effective_date == date(2026, 9, 15)
        v = p.variants.get()
        assert (v.sale_price, v.retail_price) == (Decimal("1255"), Decimal("1495"))
        assert v.packagings.get(name="carton").qty_in_base == Decimal("24")
        assert Product.objects.get(sku="62300UC").category.name == "Tinta Color"
        assert Product.objects.get(sku="12110").category.name == "Color Chameleon"
        assert Product.objects.get(sku="21500").category.name == "Dandruff Detox"  # heading row with stray zeros
        assert Product.objects.get(sku="21500").cost_price == Decimal("3001.5")
        timer = Product.objects.get(sku="53560")
        assert timer.packing == 1 and not timer.variants.get().packagings.filter(name="carton").exists()
        assert set(Category.objects.values_list("name", flat=True)) == {"Tinta Color", "Color Chameleon", "Dandruff Detox", "Accessories"}


def test_reimport_updates_prices_in_place(api, tenant):
    from apps.catalog.models import Product

    run_import(api, price_list_xlsx())
    job = run_import(api, price_list_xlsx(prices=("1200", "1333", "1590")))
    assert job["summary"] == {**job["summary"], "created": 0, "updated": 5, "categories_created": 0}
    with tenant_context(tenant, tenant.owner):
        assert Product.objects.filter(sku="62001").count() == 1
        p = Product.objects.get(sku="62001")
        assert (p.cost_price, p.sale_price, p.retail_price) == (Decimal("1200"), Decimal("1333"), Decimal("1590"))
        assert p.variants.get().sale_price == Decimal("1333")  # default variant follows the product price


def test_flat_csv_template_still_works(api, tenant):
    from apps.catalog.models import Product

    csv = "Code,Product Name,Category,Brand,UNIT,Packing,D.P,T.P,R.P,Barcode\nA-1,Alpha Shampoo,Care,Acme,300ml,12,100,120,150,1234567890123\n"
    job = run_import(api, csv.encode(), name="products.csv")
    assert job["status"] == "done" and job["rows_ok"] == 1, job["errors"]
    with tenant_context(tenant, tenant.owner):
        p = Product.objects.get(sku="A-1")
        assert p.category.name == "Care" and p.brand.name == "Acme" and p.retail_price == Decimal("150")
        assert p.variants.get().barcodes.get().code == "1234567890123"


def test_product_edit_propagates_prices_and_carton(api, tenant):
    from apps.catalog.models import Product

    run_import(api, price_list_xlsx())
    with tenant_context(tenant, tenant.owner):
        pid = Product.objects.get(sku="62001").id
    r = api.patch(f"/api/v1/products/{pid}", {"sale_price": "1300", "packing": 12}, format="json")
    assert r.status_code == 200, r.content
    v = r.data["variants"][0]
    assert Decimal(v["sale_price"]) == Decimal("1300")
    assert [Decimal(k["qty_in_base"]) for k in v["packagings"] if k["name"] == "carton"] == [Decimal("12")]


def order_sheet_xlsx():
    """Supplier order sheet: same layout as the price list but with an unlabeled quantity column at the right."""
    wb = openpyxl.Workbook()
    ws = wb.active
    for r in [
        ["Keune Pakistan"],
        ["Lahore Distribution Order"],
        [],
        ["TINTA COLOR"],
        ["S.No", "Code", "Product Name", None, None, None, None, None, "UNIT", None, None],
        [1, 62001, "NO. 1 BLACK - 60 ML", None, None, None, None, None, "60ml", None, 2149],
        [2, "62300UC", "NO. 3.00 DARK BROWN (UC) - 60 ML", None, None, None, None, None, "60ml", None, None],
        [None, None, "COLOR CHAMELEON"],
        [3, 12110, "COLOR CHAMELEON YELLOW", None, None, None, None, None, "60ml", None, 24],
        [4, 77777, "NOT IN CATALOGUE", None, None, None, None, None, "60ml", None, 5],
    ]:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_order_sheet_creates_draft_po_with_attachment(api, tenant):
    from apps.core.models import Attachment
    from apps.inventory.models import Warehouse
    from apps.partners.models import Supplier
    from apps.purchasing.models import PurchaseOrder

    with tenant_context(tenant, tenant.owner):
        Warehouse.objects.create(code="MAIN", name="Main", type="main")
        sup = Supplier.objects.create(code="S-1", name="Keune Pakistan")
    run_import(api, price_list_xlsx(), {"supplier": str(sup.id)})
    data = {"type": "purchase_order", "file": SimpleUploadedFile("LHR-Order-20-9-26.xlsx", order_sheet_xlsx()), "options": '{"filename": "LHR-Order-20-9-26.xlsx"}'}
    r = api.post("/api/v1/imports", data, format="multipart")
    assert r.status_code == 201, r.content
    job = r.data
    assert job["status"] == "done" and job["rows_total"] == 3 and job["rows_ok"] == 2 and job["rows_failed"] == 1, job
    assert "77777" in job["errors"][0]["error"]
    s = job["summary"]
    assert s["lines"] == 2 and s["skipped_no_qty"] == 1 and s["po_number"] and s["reference"] == "Lahore Distribution Order" and s["date"] == "2026-09-20"
    with tenant_context(tenant, tenant.owner):
        po = PurchaseOrder.objects.get(pk=s["po_id"])
        assert po.status == "draft" and po.supplier == sup and po.reference == "Lahore Distribution Order" and str(po.date) == "2026-09-20"
        lines = {l.variant.sku: l for l in po.lines.all()}
        assert lines["62001"].qty_ordered == Decimal("2149") and lines["62001"].unit_price == Decimal("1130")  # D.P from the catalogue
        assert lines["12110"].qty_ordered == Decimal("24")
        assert po.subtotal == Decimal("2149") * Decimal("1130") + Decimal("24") * Decimal("1431")
        att = Attachment.objects.get(model="PurchaseOrder", object_id=str(po.id))
        assert att.name == "LHR-Order-20-9-26.xlsx" and att.size > 0
    r = api.get(f"/api/v1/purchase-orders/{s['po_id']}")
    assert r.status_code == 200 and r.data["reference"] == "Lahore Distribution Order" and r.data["lines"][0]["packing"] == 24


def test_order_sheet_in_cartons(api, tenant):
    from apps.inventory.models import Warehouse
    from apps.partners.models import Supplier
    from apps.purchasing.models import PurchaseOrder

    with tenant_context(tenant, tenant.owner):
        Warehouse.objects.create(code="MAIN", name="Main", type="main")
        sup = Supplier.objects.create(code="S-1", name="Keune Pakistan")
    run_import(api, price_list_xlsx(), {"supplier": str(sup.id)})
    csv = "Code,Product Name,Cartons\n62001,NO. 1 BLACK - 60 ML,2\n"
    data = {"type": "purchase_order", "file": SimpleUploadedFile("order.csv", csv.encode()), "options": '{"supplier": "Keune Pakistan", "reference": "Test cartons"}'}
    r = api.post("/api/v1/imports", data, format="multipart")
    assert r.status_code == 201 and r.data["status"] == "done", r.content
    with tenant_context(tenant, tenant.owner):
        po = PurchaseOrder.objects.get(pk=r.data["summary"]["po_id"])
        line = po.lines.get()
        assert line.qty_ordered == Decimal("48") and line.packaging.name == "carton" and po.reference == "Test cartons"


def test_import_sets_location_and_opening_stock_once(api, tenant):
    from apps.catalog.models import Product
    from apps.inventory.models import Bin, StockMovement, Warehouse

    with tenant_context(tenant, tenant.owner):
        Warehouse.objects.create(code="MAIN", name="Main", type="main")
    csv = "Code,Product Name,UNIT,Packing,D.P,T.P,R.P,Location,Opening Qty\nX-1,Alpha,60ml,24,100,120,150,A-01,48\nX-2,Beta,60ml,12,50,60,75,A-01,\n"
    job = run_import(api, csv.encode(), name="products.csv")
    assert job["status"] == "done" and job["rows_ok"] == 2, job["errors"]
    assert job["summary"]["locations_created"] == 1 and job["summary"]["opening_stock_posted"] == 1
    with tenant_context(tenant, tenant.owner):
        bin_ = Bin.objects.get(code="A-01")
        assert {p.sku: p.location_id for p in Product.objects.all()} == {"X-1": bin_.id, "X-2": bin_.id}
        mv = StockMovement.objects.get(movement_type="opening")
        assert mv.qty == Decimal("48") and mv.bin_id == bin_.id and mv.unit_cost == Decimal("100")
    # a second upload of the same sheet updates prices but must not double the opening stock
    job2 = run_import(api, csv.encode(), name="products.csv")
    assert job2["summary"]["opening_stock_posted"] == 0 and job2["summary"]["locations_created"] == 0
    with tenant_context(tenant, tenant.owner):
        assert StockMovement.objects.filter(movement_type="opening").count() == 1


def supplier_sales_order_xlsx(so_number="51145"):
    """The supplier's own 'SALES ORDER' print: labelled header (S.O. NUMBER / DATE) and Item / Item Description / Qty / Rate."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sales_Order_Print"
    for r in [
        [None, "Keune Pakistan"],
        [None, "SALES ORDER"],
        [],
        [None, "CUSTOMER NAME : 1000073 - RO LHR", None, None, None, None, None, None, None, None, None, None, None, "S.O. NUMBER:", None, None, so_number],
        [None, None, None, None, None, None, None, None, None, None, None, None, None, "S.O. DATE:", None, None, "21-Sep-2026"],
        [],
        [None, "S.No.", "Item ", None, "Item Description", None, None, None, None, "Unit", None, None, "Qty", None, "Rate", None, "Amount"],
        [None, 1, 62001, None, "NO. 1 BLACK - 60 ML", None, None, None, None, "60ml", None, None, 288, None, 1130, None, 325440],
        [None, 2, "62300UC", None, "NO. 3.00 DARK BROWN (UC) - 60 ML", None, None, None, None, "60ml", None, None, 24, None, 1166, None, 27984],
        [None, 3, 99999, None, "NOT IN CATALOGUE", None, None, None, None, "60ml", None, None, 6, None, 500, None, 3000],
        [None, "T O T A L :", None, None, None, None, None, None, None, None, None, None, 318, None, None, None, 356424],
    ]:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_supplier_sales_order_print_becomes_draft_po(api, tenant):
    from apps.inventory.models import Warehouse
    from apps.partners.models import Supplier
    from apps.purchasing.models import PurchaseOrder

    with tenant_context(tenant, tenant.owner):
        Warehouse.objects.create(code="MAIN", name="Main", type="main")
        sup = Supplier.objects.create(code="S-1", name="Keune Pakistan")
    run_import(api, price_list_xlsx(), {"supplier": str(sup.id)})
    data = {"type": "purchase_order", "file": SimpleUploadedFile("51145 Lahore 21-09-26.xlsx", supplier_sales_order_xlsx()), "options": '{"filename": "51145 Lahore 21-09-26.xlsx"}'}
    r = api.post("/api/v1/imports", data, format="multipart")
    assert r.status_code == 201, r.content
    s = r.data["summary"]
    assert r.data["rows_ok"] == 2 and r.data["rows_failed"] == 1 and "99999" in r.data["errors"][0]["error"]
    assert s["supplier_ref"] == "51145" and s["date"] == "2026-09-21" and s["reference"].startswith("S.O. 51145")
    with tenant_context(tenant, tenant.owner):
        po = PurchaseOrder.objects.get(pk=s["po_id"])
        assert po.status == "draft" and po.supplier_ref == "51145" and str(po.date) == "2026-09-21"
        lines = {l.variant.sku: l for l in po.lines.all()}
        assert lines["62001"].qty_ordered == Decimal("288") and lines["62001"].unit_price == Decimal("1130")  # rate from the sheet
        assert lines["62300UC"].unit_price == Decimal("1166")
        assert po.subtotal == Decimal("288") * Decimal("1130") + Decimal("24") * Decimal("1166")

    # the same supplier document must not slip in twice
    again = api.post("/api/v1/imports", {"type": "purchase_order", "file": SimpleUploadedFile("51145.xlsx", supplier_sales_order_xlsx()), "options": '{"filename": "51145.xlsx"}'}, format="multipart")
    assert again.data["status"] == "failed" and "already imported" in again.data["errors"][0]["error"]
    ok = api.post("/api/v1/imports", {"type": "purchase_order", "file": SimpleUploadedFile("51145.xlsx", supplier_sales_order_xlsx()), "options": '{"filename": "51145.xlsx", "allow_duplicate": true}'}, format="multipart")
    assert ok.data["summary"]["po_number"] and ok.data["rows_ok"] == 2


def test_import_fills_the_item_code_and_the_barcodes(api, tenant, tmp_path):
    """
    A price list's "Code" column is the supplier's item code, and a sheet may also carry the unit and box barcodes.
    After an import the product is scannable and the code people quote is on it.
    """
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Code", "Product Name", "UNIT", "Packing", "D.P", "T.P", "R.P", "UPC", "Box UPC"])
    ws.append(["77001", "IMPORTED SHAMPOO 1000ML", "1000ml", 6, 1000, 1200, 1500, "8964777000016", "18964777000013"])
    path = tmp_path / "with-codes.xlsx"
    wb.save(path)

    with open(path, "rb") as fh:
        r = api.post("/api/v1/imports", {"type": "products", "file": fh, "options": '{"update_existing": true}'}, format="multipart")
    assert r.status_code in (200, 201), r.content
    assert r.data["summary"]["created"] == 1, r.data

    found = api.get("/api/v1/products", {"search": "77001"})
    product = found.data["results"][0]
    assert product["item_code"] == "77001"
    assert product["upc"] == "8964777000016"
    assert product["box_upc"] == "18964777000013"
    assert product["packing"] == 6

    # both codes resolve, and the box one counts a whole box
    unit = api.get("/api/v1/barcodes/resolve", {"code": "8964777000016"})
    assert unit.status_code == 200
    from apps.core.tenancy import tenant_context

    with tenant_context(tenant, tenant.owner):
        from apps.catalog.models import Barcode

        assert Barcode.objects.get(code="18964777000013").packaging.qty_in_base == 6
